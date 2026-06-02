"""
DOCX report generator using python-docx.
Produces two report types:
  • generate_docx_report()      – Full penetration test report
  • generate_remediation_docx() – Remediation roadmap
Both include the Bouquet Innovation Lda logo in the header.
"""

import os, io
from datetime import datetime

_BASE     = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
LOGO_PATH = os.path.join(_BASE, "static", "img", "bouquet_logo.png")

try:
    from docx import Document
    from docx.shared import Inches, Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    import docx.opc.constants
    DOCX_OK = True
except ImportError:
    DOCX_OK = False

# ── Colours ──────────────────────────────────────────────────────────────────
GOLD      = RGBColor(0xB4, 0x8C, 0x14)
GOLD_DARK = RGBColor(0x7A, 0x5E, 0x00)
BLACK     = RGBColor(0x1A, 0x1A, 0x1A)
GREY      = RGBColor(0x55, 0x55, 0x55)
WHITE_RGB = RGBColor(0xFF, 0xFF, 0xFF)

SEV_COLORS = {
    "CRITICAL": (RGBColor(0xFF,0xFF,0xFF), RGBColor(0x7D,0x11,0x28), "7D1128"),
    "HIGH":     (RGBColor(0xFF,0xFF,0xFF), RGBColor(0x9B,0x23,0x35), "9B2335"),
    "MEDIUM":   (RGBColor(0x1A,0x1A,0x1A), RGBColor(0xFF,0xE8,0xB0), "7A4F00"),
    "LOW":      (RGBColor(0x1A,0x1A,0x1A), RGBColor(0xC8,0xED,0xD4), "1A5C2A"),
    "INFO":     (RGBColor(0x1A,0x1A,0x1A), RGBColor(0xC8,0xD9,0xF0), "1A3A5C"),
}

# ── Helpers ───────────────────────────────────────────────────────────────────

def _set_cell_bg(cell, hex_color: str):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex_color.lstrip("#"))
    tcPr.append(shd)


def _set_cell_border(cell, **edges):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top","left","bottom","right","insideH","insideV"):
        if edge in edges:
            tag = OxmlElement(f"w:{edge}")
            tag.set(qn("w:val"),   edges[edge].get("val","single"))
            tag.set(qn("w:sz"),    str(edges[edge].get("sz", 4)))
            tag.set(qn("w:color"), edges[edge].get("color","000000"))
            tcBorders.append(tag)
    tcPr.append(tcBorders)


def _bold_run(para, text, size=10, color=None):
    run = para.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return run


def _normal_run(para, text, size=9, color=None, italic=False):
    run = para.add_run(text)
    run.font.size = Pt(size)
    run.italic = italic
    if color:
        run.font.color.rgb = color
    return run


def _add_header(doc, meta, report_type):
    """Document header with logo, title, metadata table."""
    # Logo
    if os.path.exists(LOGO_PATH):
        doc.add_picture(LOGO_PATH, width=Inches(1.1))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Gold divider
    p = doc.add_paragraph()
    p.paragraph_format.space_after  = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run("━" * 80)
    run.font.color.rgb = GOLD
    run.font.size = Pt(7)

    # Company name
    p = doc.add_paragraph()
    r = p.add_run("BOUQUET INNOVATION Lda  |  Cybersecurity & Technology Consultancy")
    r.font.size = Pt(9); r.font.color.rgb = GOLD_DARK; r.bold = True
    p.paragraph_format.space_after = Pt(4)

    # Report type
    p = doc.add_paragraph()
    r = p.add_run(report_type)
    r.font.size = Pt(11); r.font.color.rgb = GREY

    # Title
    p = doc.add_paragraph()
    r = p.add_run(meta.get("title","Penetration Test Report"))
    r.font.size = Pt(22); r.bold = True; r.font.color.rgb = GOLD_DARK
    p.paragraph_format.space_after = Pt(8)

    # Metadata table
    fields = [
        ("Target",          meta.get("target","—")),
        ("Engagement Type", meta.get("engagement_type","—")),
        ("Testing Period",  meta.get("testing_period","—")),
        ("Tester",          meta.get("tester","—")),
        ("Date",            meta.get("date","—")),
        ("Classification",  meta.get("classification","CONFIDENTIAL")),
    ]
    tbl = doc.add_table(rows=len(fields), cols=2)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, (key, val) in enumerate(fields):
        row = tbl.rows[i]
        row.cells[0].width = Cm(4.5)
        row.cells[1].width = Cm(12)
        _set_cell_bg(row.cells[0], "F5F5F5")
        p0 = row.cells[0].paragraphs[0]
        r0 = p0.add_run(key)
        r0.bold = True; r0.font.size = Pt(8); r0.font.color.rgb = GOLD_DARK
        p1 = row.cells[1].paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(8)
    doc.add_paragraph()  # spacer

    # Disclaimer
    p = doc.add_paragraph()
    r = p.add_run(
        "This document is CONFIDENTIAL and intended solely for the named recipient. "
        "It contains sensitive security information. Distribution to unauthorised persons is prohibited. "
        "Prepared by Bouquet Innovation Lda using Wouapit-Hack.")
    r.italic = True; r.font.size = Pt(8); r.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(12)


def _section(doc, number, title):
    p = doc.add_paragraph()
    run = p.add_run(f"{number}. {title}")
    run.bold = True; run.font.size = Pt(13)
    run.font.color.rgb = GOLD_DARK
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after  = Pt(4)
    # Gold underline
    p2 = doc.add_paragraph()
    r2  = p2.add_run("─" * 90)
    r2.font.size = Pt(6); r2.font.color.rgb = GOLD
    p2.paragraph_format.space_after = Pt(4)


def _subsection(doc, title):
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.bold = True; r.font.size = Pt(10); r.font.color.rgb = BLACK
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after  = Pt(3)


def _sev_summary_table(doc, findings):
    sev_order = ["CRITICAL","HIGH","MEDIUM","LOW","INFO"]
    counts    = {s: sum(1 for f in findings if f.get("severity","INFO").upper()==s)
                 for s in sev_order}
    slas      = {"CRITICAL":"24 hours","HIGH":"72 hours","MEDIUM":"30 days","LOW":"90 days","INFO":"Advisory"}

    tbl = doc.add_table(rows=1+len(sev_order), cols=4)
    tbl.style = "Table Grid"

    # Header
    hdrs = ["SEVERITY","COUNT","REMEDIATION SLA","RISK LEVEL"]
    risk = {"CRITICAL":"Critical – Act immediately","HIGH":"High – Act within 72 h",
            "MEDIUM":"Medium – 30 days","LOW":"Low – Schedule","INFO":"Informational"}
    for j, h in enumerate(hdrs):
        cell = tbl.rows[0].cells[j]
        _set_cell_bg(cell, "F5F5F5")
        p = cell.paragraphs[0]; r = p.add_run(h)
        r.bold = True; r.font.size = Pt(8); r.font.color.rgb = GOLD_DARK

    for i, sev in enumerate(sev_order, 1):
        row   = tbl.rows[i]
        tc, bg_c, hex_c = SEV_COLORS.get(sev, (BLACK, WHITE_RGB, "FFFFFF"))
        # Severity cell
        _set_cell_bg(row.cells[0], hex_c)
        _set_cell_bg(row.cells[1], hex_c)
        p = row.cells[0].paragraphs[0]; r = p.add_run(sev)
        r.bold = True; r.font.size = Pt(9); r.font.color.rgb = tc
        # Count
        p2 = row.cells[1].paragraphs[0]; p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(str(counts[sev]))
        r2.bold = True; r2.font.size = Pt(14); r2.font.color.rgb = tc
        # SLA
        p3 = row.cells[2].paragraphs[0]; r3 = p3.add_run(slas[sev])
        r3.font.size = Pt(8)
        # Risk
        p4 = row.cells[3].paragraphs[0]; r4 = p4.add_run(risk[sev])
        r4.font.size = Pt(8)
    doc.add_paragraph()


def _finding_block_docx(doc, f, idx, remediation_only=False):
    sev = f.get("severity","INFO").upper()
    tc, bg_c, hex_c = SEV_COLORS.get(sev, (BLACK, WHITE_RGB, "FFFFFF"))

    # Title bar
    tbl = doc.add_table(rows=1, cols=2)
    tbl.style = "Table Grid"
    tbl.columns[0].width = Cm(14)
    tbl.columns[1].width = Cm(2.8)
    _set_cell_bg(tbl.rows[0].cells[0], hex_c)
    _set_cell_bg(tbl.rows[0].cells[1], hex_c)
    p = tbl.rows[0].cells[0].paragraphs[0]
    r = p.add_run(f"#{idx}  {f.get('title','Untitled')}")
    r.bold = True; r.font.size = Pt(10)
    r.font.color.rgb = tc if hex_c not in ("FFFFFF","F5F5F5","C8EDD4","C8D9F0","FFE8B0") else BLACK
    p2 = tbl.rows[0].cells[1].paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(sev)
    r2.bold = True; r2.font.size = Pt(8); r2.font.color.rgb = tc
    doc.add_paragraph()

    def field(label, value, is_code=False):
        if not value or str(value).strip() in ("","N/A","None"):
            return
        lp = doc.add_paragraph()
        lp.paragraph_format.space_after = Pt(1)
        lr = lp.add_run(label.upper())
        lr.bold = True; lr.font.size = Pt(7); lr.font.color.rgb = GREY
        vp = doc.add_paragraph()
        vp.paragraph_format.left_indent = Cm(0.5)
        vp.paragraph_format.space_after = Pt(4)
        vr = vp.add_run(str(value)[:1000])
        vr.font.size = Pt(8.5)
        if is_code:
            vr.font.name = "Courier New"
            vp.paragraph_format.left_indent = Cm(0.5)

    if not remediation_only:
        field("Description",              f.get("description"))
        field("Location / URL",           f.get("location"))
        field("CVSS Score",               f.get("cvss"))
        field("Evidence / Proof of Concept", f.get("evidence"), is_code=True)

    field("Remediation", f.get("remediation"))

    p = doc.add_paragraph()
    r = p.add_run("─" * 90)
    r.font.size = Pt(6); r.font.color.rgb = RGBColor(0xDD,0xDD,0xDD)
    p.paragraph_format.space_after = Pt(6)


# ═══════════════════════════════════════════════════════════════════════════════
# PUBLIC: Full pentest DOCX
# ═══════════════════════════════════════════════════════════════════════════════
def generate_docx_report(data: dict, output_dir: str) -> dict:
    if not DOCX_OK:
        return {"error": "python-docx not installed. pip install python-docx"}

    os.makedirs(output_dir, exist_ok=True)
    now      = datetime.utcnow()
    findings = data.get("findings", [])
    meta = {
        "title":           data.get("title", "Penetration Test Report"),
        "target":          data.get("target", "Unknown"),
        "tester":          data.get("tester", "Bouquet Innovation Lda"),
        "date":            now.strftime("%d %B %Y"),
        "engagement_type": data.get("engagement_type", "Black Box"),
        "classification":  data.get("classification", "CONFIDENTIAL"),
        "testing_period":  data.get("testing_period", now.strftime("%B %Y")),
        "tools_used":      data.get("tools_used", "Wouapit-Hack"),
    }

    slug     = meta["target"].replace(".", "_").replace("/", "_").replace(":", "_")
    filename = f"pentest_{slug}_{now.strftime('%Y%m%d_%H%M%S')}.docx"
    filepath = os.path.join(output_dir, filename)

    doc = Document()
    # Page margins
    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    _add_header(doc, meta, "PENETRATION TEST REPORT")

    # 1. Executive Summary
    _section(doc, 1, "Executive Summary")
    p = doc.add_paragraph(data.get("executive_summary",
        f"A penetration test was conducted against {meta['target']}. "
        f"The assessment identified {len(findings)} finding(s)."))
    p.runs[0].font.size = Pt(9)
    doc.add_paragraph()
    _sev_summary_table(doc, findings)

    # 2. Scope & Methodology
    _section(doc, 2, "Scope & Methodology")
    rows = [
        ("Target",          meta["target"]),
        ("Engagement Type", meta["engagement_type"]),
        ("Testing Period",  meta["testing_period"]),
        ("Tester",          meta["tester"]),
        ("Tools Used",      meta["tools_used"]),
        ("Methodology",     "OWASP Testing Guide v4, PTES, OWASP Top 10"),
    ]
    tbl = doc.add_table(rows=len(rows), cols=2)
    tbl.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        _set_cell_bg(tbl.rows[i].cells[0], "F5F5F5")
        pr = tbl.rows[i].cells[0].paragraphs[0]
        rr = pr.add_run(k); rr.bold = True; rr.font.size = Pt(8); rr.font.color.rgb = GOLD_DARK
        pv = tbl.rows[i].cells[1].paragraphs[0]
        rv = pv.add_run(v); rv.font.size = Pt(8)
    doc.add_paragraph()

    # 3. Recon Results
    recon = data.get("recon", {})
    if any(recon.get(k) for k in ("open_ports","subdomains","technologies","headers")):
        _section(doc, 3, "Reconnaissance Results")

        if recon.get("technologies"):
            _subsection(doc, "Technologies Detected")
            p = doc.add_paragraph(", ".join(recon["technologies"]))
            p.runs[0].font.size = Pt(9)

        if recon.get("open_ports"):
            _subsection(doc, "Open Ports")
            t = doc.add_table(rows=1+len(recon["open_ports"][:40]), cols=3)
            t.style = "Table Grid"
            for j, h in enumerate(["PORT","STATE","SERVICE"]):
                c = t.rows[0].cells[j]; _set_cell_bg(c, "F5F5F5")
                r2 = c.paragraphs[0].add_run(h)
                r2.bold = True; r2.font.size = Pt(8); r2.font.color.rgb = GOLD_DARK
            for i, p2 in enumerate(recon["open_ports"][:40], 1):
                bg = "F0FFF4" if i % 2 == 0 else "FFFFFF"
                for j2, val in enumerate([str(p2.get("port","")), "OPEN", p2.get("service","")]):
                    _set_cell_bg(t.rows[i].cells[j2], bg)
                    rv2 = t.rows[i].cells[j2].paragraphs[0].add_run(val)
                    rv2.font.size = Pt(8)
                    if j2 == 1:
                        rv2.font.color.rgb = RGBColor(0x1A,0x5C,0x2A); rv2.bold = True
            doc.add_paragraph()

        if recon.get("subdomains"):
            _subsection(doc, "Subdomains Found")
            t = doc.add_table(rows=1+len(recon["subdomains"][:30]), cols=2)
            t.style = "Table Grid"
            for j, h in enumerate(["SUBDOMAIN","IP ADDRESS"]):
                c = t.rows[0].cells[j]; _set_cell_bg(c, "F5F5F5")
                r2 = c.paragraphs[0].add_run(h)
                r2.bold = True; r2.font.size = Pt(8); r2.font.color.rgb = GOLD_DARK
            for i, s2 in enumerate(recon["subdomains"][:30], 1):
                bg = "FAFAFA" if i % 2 == 0 else "FFFFFF"
                for j2, val in enumerate([s2.get("subdomain",""), s2.get("ip","")]):
                    _set_cell_bg(t.rows[i].cells[j2], bg)
                    t.rows[i].cells[j2].paragraphs[0].add_run(val).font.size = Pt(8)
            doc.add_paragraph()

        if recon.get("headers"):
            _subsection(doc, "Security Headers Audit")
            items2 = list(recon["headers"].items())
            t = doc.add_table(rows=1+len(items2), cols=2)
            t.style = "Table Grid"
            for j, h in enumerate(["HEADER","STATUS"]):
                c = t.rows[0].cells[j]; _set_cell_bg(c, "F5F5F5")
                r2 = c.paragraphs[0].add_run(h)
                r2.bold = True; r2.font.size = Pt(8); r2.font.color.rgb = GOLD_DARK
            for i, (k2, v2) in enumerate(items2, 1):
                bg = "FFF0F0" if v2 == "MISSING" else "FFFFFF"
                _set_cell_bg(t.rows[i].cells[0], bg)
                _set_cell_bg(t.rows[i].cells[1], bg)
                t.rows[i].cells[0].paragraphs[0].add_run(k2).font.size = Pt(8)
                r3 = t.rows[i].cells[1].paragraphs[0].add_run(
                    ("✗ MISSING" if v2 == "MISSING" else "✓ "+str(v2)[:60]))
                r3.font.size = Pt(8)
                r3.font.color.rgb = (RGBColor(0x7D,0x11,0x28) if v2 == "MISSING"
                                     else RGBColor(0x1A,0x5C,0x2A))
            doc.add_paragraph()

    # 4. Detailed Findings
    doc.add_page_break()
    _section(doc, 4, "Detailed Findings")
    if not findings:
        doc.add_paragraph("No vulnerabilities were identified.").runs[0].font.size = Pt(9)
    else:
        for i, f in enumerate(findings, 1):
            _finding_block_docx(doc, f, i)

    # 5. Recommendations
    doc.add_page_break()
    _section(doc, 5, "Recommendations")
    for rec in [
        "Immediately remediate all Critical and High severity findings.",
        "Apply the principle of least privilege across all systems and accounts.",
        "Implement and enforce a Web Application Firewall (WAF).",
        "Enable Multi-Factor Authentication on all externally accessible accounts.",
        "Enforce HTTPS everywhere and configure HSTS.",
        "Add all missing security headers to web server configurations.",
        "Conduct regular penetration tests (minimum annually).",
        "Keep all software components, libraries and OS packages patched.",
        "Perform security awareness training for all staff quarterly.",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(rec); r.font.size = Pt(9)

    # Footer note
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(
        f"Prepared by Bouquet Innovation Lda using Wouapit-Hack  |  {meta['date']}  |  CONFIDENTIAL")
    r.italic = True; r.font.size = Pt(8); r.font.color.rgb = GREY

    doc.save(filepath)
    counts = {s: sum(1 for f in findings if f.get("severity","INFO").upper()==s)
              for s in ("CRITICAL","HIGH","MEDIUM","LOW","INFO")}
    return {"success": True, "filename": filename, "filepath": filepath, "type": "pentest_docx",
            "summary": {"total_findings": len(findings), "severity_counts": counts}}


# ═══════════════════════════════════════════════════════════════════════════════
# PUBLIC: Remediation DOCX
# ═══════════════════════════════════════════════════════════════════════════════
def generate_remediation_docx(data: dict, output_dir: str) -> dict:
    if not DOCX_OK:
        return {"error": "python-docx not installed. pip install python-docx"}

    os.makedirs(output_dir, exist_ok=True)
    now      = datetime.utcnow()
    findings = data.get("findings", [])
    meta = {
        "title":           data.get("title", "Remediation Report"),
        "target":          data.get("target", "Unknown"),
        "tester":          data.get("tester", "Bouquet Innovation Lda"),
        "date":            now.strftime("%d %B %Y"),
        "engagement_type": data.get("engagement_type", "—"),
        "classification":  data.get("classification", "CONFIDENTIAL"),
        "testing_period":  data.get("testing_period", now.strftime("%B %Y")),
    }

    slug     = meta["target"].replace(".", "_").replace("/", "_").replace(":", "_")
    filename = f"remediation_{slug}_{now.strftime('%Y%m%d_%H%M%S')}.docx"
    filepath = os.path.join(output_dir, filename)

    doc = Document()
    for section in doc.sections:
        section.top_margin    = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    _add_header(doc, meta, "REMEDIATION REPORT")

    sev_order = ["CRITICAL","HIGH","MEDIUM","LOW","INFO"]
    by_sev    = {s: [f for f in findings if f.get("severity","INFO").upper()==s]
                 for s in sev_order}
    slas  = {"CRITICAL":"24 hours","HIGH":"72 hours","MEDIUM":"30 days","LOW":"90 days","INFO":"Advisory"}
    effort= {"CRITICAL":"High","HIGH":"High","MEDIUM":"Medium","LOW":"Low","INFO":"Minimal"}

    # 1. Purpose
    _section(doc, 1, "Purpose")
    p = doc.add_paragraph(
        "This report provides a prioritised, actionable remediation roadmap for all security findings. "
        "Address Critical items immediately before moving to lower-priority issues.")
    p.runs[0].font.size = Pt(9)

    # 2. Priority Matrix
    _section(doc, 2, "Remediation Priority Matrix")
    tbl = doc.add_table(rows=1+len(sev_order), cols=5)
    tbl.style = "Table Grid"
    for j, h in enumerate(["SEVERITY","COUNT","SLA","EFFORT","PRIORITY"]):
        c = tbl.rows[0].cells[j]; _set_cell_bg(c, "F5F5F5")
        r2 = c.paragraphs[0].add_run(h)
        r2.bold = True; r2.font.size = Pt(8); r2.font.color.rgb = GOLD_DARK
    p_labels = {"CRITICAL":"P1","HIGH":"P2","MEDIUM":"P3","LOW":"P4","INFO":"P5"}
    for i, sev in enumerate(sev_order, 1):
        tc, _, hex_c = SEV_COLORS.get(sev,(BLACK,WHITE_RGB,"FFFFFF"))
        row = tbl.rows[i]
        _set_cell_bg(row.cells[0], hex_c)
        _set_cell_bg(row.cells[1], hex_c)
        r3 = row.cells[0].paragraphs[0].add_run(sev)
        r3.bold = True; r3.font.size = Pt(9); r3.font.color.rgb = tc
        r4 = row.cells[1].paragraphs[0].add_run(str(len(by_sev[sev])))
        r4.bold = True; r4.font.size = Pt(13); r4.font.color.rgb = tc
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for j2, val in enumerate([slas[sev], effort[sev], p_labels[sev]], 2):
            r5 = row.cells[j2].paragraphs[0].add_run(val)
            r5.font.size = Pt(8)
    doc.add_paragraph()

    # 3. Roadmap
    _section(doc, 3, "Remediation Roadmap")
    phases = [("Phase 1 — Immediate (0–24 h)","CRITICAL"),
              ("Phase 2 — Short Term (1–3 days)","HIGH"),
              ("Phase 3 — Medium Term (30 days)","MEDIUM"),
              ("Phase 4 — Long Term (90 days)","LOW"),
              ("Phase 5 — Advisory","INFO")]
    for label, sev in phases:
        items = by_sev[sev]
        if not items:
            continue
        tc, _, hex_c = SEV_COLORS.get(sev,(BLACK,WHITE_RGB,"FFFFFF"))
        pt = doc.add_table(rows=1, cols=1)
        pt.style = "Table Grid"
        _set_cell_bg(pt.rows[0].cells[0], hex_c)
        r6 = pt.rows[0].cells[0].paragraphs[0].add_run(
            f"{label}  ({len(items)} item{'s' if len(items)!=1 else ''})")
        r6.bold = True; r6.font.size = Pt(10); r6.font.color.rgb = tc
        doc.add_paragraph()
        for item in items:
            p7 = doc.add_paragraph(style="List Bullet")
            r7 = p7.add_run(f"{item.get('title','')}  ")
            r7.bold = True; r7.font.size = Pt(9); r7.font.color.rgb = GOLD_DARK
            r8 = p7.add_run(f"[CVSS {item.get('cvss','N/A')} | {str(item.get('location',''))[:50]}]")
            r8.font.size = Pt(8); r8.font.color.rgb = GREY
        doc.add_paragraph()

    # 4. Detailed Steps
    doc.add_page_break()
    _section(doc, 4, "Detailed Remediation Steps")
    idx = 1
    for sev in sev_order:
        items = by_sev[sev]
        if not items:
            continue
        _subsection(doc, f"{sev} Severity Findings")
        for f in items:
            _finding_block_docx(doc, f, idx, remediation_only=True)
            idx += 1

    # 5. Checklist
    doc.add_page_break()
    _section(doc, 5, "Post-Remediation Verification Checklist")
    checks = [
        ("All Critical findings patched or mitigated", "CRITICAL"),
        ("All High findings patched or mitigated", "HIGH"),
        ("Security headers configured on web server", "MEDIUM"),
        ("SSL/TLS certificates valid and up to date", "MEDIUM"),
        ("Unnecessary ports firewalled", "MEDIUM"),
        ("Default credentials changed on all services", "HIGH"),
        ("Input validation on all user-supplied inputs", "CRITICAL"),
        ("Error messages do not expose stack traces", "MEDIUM"),
        ("Patch management process documented", "LOW"),
        ("Security training completed for dev team", "LOW"),
        ("Re-test scheduled with security team", "INFO"),
    ]
    tbl2 = doc.add_table(rows=1+len(checks), cols=5)
    tbl2.style = "Table Grid"
    for j, h in enumerate(["☐","TASK","SEV","OWNER","DATE DONE"]):
        c = tbl2.rows[0].cells[j]; _set_cell_bg(c, "F5F5F5")
        r9 = c.paragraphs[0].add_run(h)
        r9.bold = True; r9.font.size = Pt(8); r9.font.color.rgb = GOLD_DARK
    for i, (task, sev) in enumerate(checks, 1):
        tc, bg_c, hex_c = SEV_COLORS.get(sev,(BLACK,WHITE_RGB,"FFFFFF"))
        bg = "FAFAFA" if i % 2 == 0 else "FFFFFF"
        tbl2.rows[i].cells[0].paragraphs[0].add_run("☐").font.size = Pt(10)
        r10 = tbl2.rows[i].cells[1].paragraphs[0].add_run(task)
        r10.font.size = Pt(8)
        c2 = tbl2.rows[i].cells[2]; _set_cell_bg(c2, hex_c)
        r11 = c2.paragraphs[0].add_run(sev)
        r11.bold = True; r11.font.size = Pt(7); r11.font.color.rgb = tc
        for j2 in [3, 4]:
            _set_cell_bg(tbl2.rows[i].cells[j2], bg)
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run(
        f"Prepared by Bouquet Innovation Lda using Wouapit-Hack  |  {meta['date']}  |  CONFIDENTIAL")
    r.italic = True; r.font.size = Pt(8); r.font.color.rgb = GREY

    doc.save(filepath)
    return {"success": True, "filename": filename, "filepath": filepath, "type": "remediation_docx",
            "summary": {"total": len(findings),
                        "phases": {s: len(by_sev[s]) for s in sev_order}}}
